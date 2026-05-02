"""LocalXlsxStorage — 로컬 엑셀(.xlsx) 백엔드.

CSV 백엔드와 동일한 BaseStorage 인터페이스를 구현하지만, 파일 형식이 .xlsx 라
사용자가 Excel/Numbers 에서 직접 열어 편집할 수 있다.
한 워크북에 두 시트(keyword_pool / content_db)를 두는 구조 — 기획서의 1-tab-per-domain 모델 그대로.

장점:
- Excel/Numbers 에서 더블클릭으로 즉시 열림
- 컬럼 타입(숫자/문자) 보존 — CSV 의 숫자가 텍스트로 변하는 문제 없음
- 비개발자 동료에게 그대로 공유 가능

단점:
- 동시 편집 어려움 → 클라우드 협업이 필요하면 Mirror + Google Sheets 조합 사용
"""
from __future__ import annotations

import os
import threading
from typing import List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from ..models import ContentRecord, KeywordPoolItem
from .base import BaseStorage


class LocalXlsxStorage(BaseStorage):
    name = "xlsx"

    POOL_SHEET = "keyword_pool"
    CONTENT_SHEET = "content_db"

    def __init__(self, data_dir: str = "./data", filename: str = "osmu_workbook.xlsx"):
        self.data_dir = data_dir
        os.makedirs(self.data_dir, exist_ok=True)
        self.path = os.path.join(self.data_dir, filename)
        # 동일 프로세스에서 동시 쓰기를 막기 위한 락
        self._lock = threading.Lock()
        self._ensure_workbook()

    # ── 내부 헬퍼 ─────────────────────────────────────
    def _ensure_workbook(self) -> None:
        if not os.path.isfile(self.path) or os.path.getsize(self.path) == 0:
            wb = Workbook()
            # 기본 시트 제거 후 두 시트 생성
            default = wb.active
            wb.remove(default)
            ws_pool = wb.create_sheet(self.POOL_SHEET)
            ws_pool.append(KeywordPoolItem.HEADER)
            ws_content = wb.create_sheet(self.CONTENT_SHEET)
            ws_content.append(ContentRecord.HEADER)
            self._auto_fit(ws_pool, KeywordPoolItem.HEADER)
            self._auto_fit(ws_content, ContentRecord.HEADER)
            wb.save(self.path)
            return
        # 기존 파일이 있는데 시트가 없으면 보완
        wb = load_workbook(self.path)
        changed = False
        if self.POOL_SHEET not in wb.sheetnames:
            ws = wb.create_sheet(self.POOL_SHEET)
            ws.append(KeywordPoolItem.HEADER)
            self._auto_fit(ws, KeywordPoolItem.HEADER)
            changed = True
        if self.CONTENT_SHEET not in wb.sheetnames:
            ws = wb.create_sheet(self.CONTENT_SHEET)
            ws.append(ContentRecord.HEADER)
            self._auto_fit(ws, ContentRecord.HEADER)
            changed = True
        if changed:
            wb.save(self.path)

    @staticmethod
    def _auto_fit(ws, header: list) -> None:
        for i, name in enumerate(header, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(12, len(str(name)) + 4)

    def _load(self):
        return load_workbook(self.path)

    def _read_rows(self, sheet_name: str) -> List[list]:
        wb = self._load()
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            # 모든 셀이 None 이면 스킵
            if all(c is None or c == "" for c in row):
                continue
            rows.append(["" if c is None else c for c in row])
        return rows

    def _write_all(self, sheet_name: str, header: list, rows: List[list]) -> None:
        with self._lock:
            wb = self._load()
            if sheet_name in wb.sheetnames:
                del wb[sheet_name]
            ws = wb.create_sheet(sheet_name)
            ws.append(header)
            for r in rows:
                # 빈 문자열은 그대로, 숫자는 숫자형으로 보존
                ws.append(r)
            self._auto_fit(ws, header)
            # 시트 순서 정리: pool, content_db 순으로
            order = [self.POOL_SHEET, self.CONTENT_SHEET]
            for i, name in enumerate(order):
                if name in wb.sheetnames:
                    wb.move_sheet(name, offset=i - wb.sheetnames.index(name))
            wb.save(self.path)

    def _append_row(self, sheet_name: str, row: list) -> None:
        with self._lock:
            wb = self._load()
            ws = wb[sheet_name]
            ws.append(row)
            wb.save(self.path)

    # ── BaseStorage: keyword_pool ──────────────────────
    def list_pool(self) -> List[KeywordPoolItem]:
        return [KeywordPoolItem.from_row(r) for r in self._read_rows(self.POOL_SHEET)]

    def get_pool(self, keyword_id: str) -> Optional[KeywordPoolItem]:
        for it in self.list_pool():
            if it.keyword_id == keyword_id:
                return it
        return None

    def upsert_pool(self, item: KeywordPoolItem) -> None:
        items = self.list_pool()
        for i, it in enumerate(items):
            if it.keyword_id == item.keyword_id:
                items[i] = item
                break
        else:
            items.append(item)
        self.replace_pool(items)

    def delete_pool(self, keyword_id: str) -> bool:
        items = self.list_pool()
        new_items = [it for it in items if it.keyword_id != keyword_id]
        if len(new_items) == len(items):
            return False
        self.replace_pool(new_items)
        return True

    def replace_pool(self, items: List[KeywordPoolItem]) -> None:
        self._write_all(self.POOL_SHEET, KeywordPoolItem.HEADER, [it.to_row() for it in items])

    # ── BaseStorage: content_db ────────────────────────
    def list_content(self) -> List[ContentRecord]:
        return [ContentRecord.from_row(r) for r in self._read_rows(self.CONTENT_SHEET)]

    def append_content(self, record: ContentRecord) -> None:
        self._append_row(self.CONTENT_SHEET, record.to_row())

    def replace_content(self, records) -> None:
        self._write_all(self.CONTENT_SHEET, ContentRecord.HEADER, [r.to_row() for r in records])
